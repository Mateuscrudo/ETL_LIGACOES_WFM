import shutil
import uuid
import unittest
import warnings
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook

from controle_utils import carregar_aba_controle
from controle_leitura import determinar_ultima_linha_cadastro


class LeituraControleTests(unittest.TestCase):
    def setUp(self):
        temp_root = Path(__file__).resolve().parents[1] / 'work' / 'test_leitura_controle'
        temp_root.mkdir(parents=True, exist_ok=True)
        fixture_dir = temp_root / uuid.uuid4().hex
        fixture_dir.mkdir()
        self.addCleanup(shutil.rmtree, fixture_dir)
        self.path = fixture_dir / 'controle.xlsx'

    def salvar(self, preencher, header=1):
        wb = Workbook()
        ws = wb.active
        ws.title = 'GERAL'
        if header > 1:
            ws.cell(1, 1, 'Relatorio')
        for i, nome in enumerate(['Matricula', 'Endereco', 'Numero', 'Complemento', 'INFO Hidrometro'], 1):
            ws.cell(header, i, nome)
        preencher(ws)
        wb.save(self.path)
        wb.close()

    def limite(self, header=0):
        return determinar_ultima_linha_cadastro(self.path, 'GERAL', header, {1, 2, 3, 4})

    def test_normal_preserva_dataframe(self):
        self.salvar(lambda ws: (ws.append(['AR0045', 'Rua A', 12, 'Casa 2', 'INSTALADO']),
                                ws.append(['AR0046', 'Rua B', 0, None, 'PENDENTE'])))
        esperado = pd.read_excel(self.path, sheet_name='GERAL')
        resultado, colunas, header = carregar_aba_controle(self.path)
        pd.testing.assert_frame_equal(resultado, esperado)
        self.assertEqual((header, colunas['matricula']), (0, 'Matricula'))

    def test_lacunas_e_filtro_desatualizado(self):
        def preencher(ws):
            ws['A2'], ws['B2'] = 'AR0045', 'Rua A'
            ws['A150'], ws['B150'] = 'AR0046', 'Rua B'
            ws.auto_filter.ref = 'A1:E2'
        self.salvar(preencher)
        self.assertEqual(self.limite(), 150)
        resultado, _, _ = carregar_aba_controle(self.path)
        self.assertEqual(len(resultado), 149)
        self.assertEqual(resultado.iloc[-1]['Matricula'], 'AR0046')

    def test_cadastros_parciais(self):
        for coluna in ('A', 'B', 'C', 'D'):
            with self.subTest(coluna=coluna):
                self.salvar(lambda ws: ws.__setitem__(coluna + '80', 'Dado parcial'))
                self.assertEqual(self.limite(), 80)
                resultado, _, _ = carregar_aba_controle(self.path)
                self.assertEqual(len(resultado), 79)

    def test_zero_e_false(self):
        for valor in (0, False):
            with self.subTest(valor=valor):
                self.salvar(lambda ws: ws.__setitem__('C75', valor))
                self.assertEqual(self.limite(), 75)
                resultado, _, _ = carregar_aba_controle(self.path)
                self.assertEqual(resultado.iloc[-1]['Numero'], valor)

    def test_formula_sem_cache_define_limite(self):
        self.salvar(lambda ws: ws.__setitem__('A90', '=1+1'))
        self.assertEqual(self.limite(), 90)
        original = self.path.read_bytes()
        real_read = pd.read_excel
        with patch('controle_utils.pd.read_excel', wraps=real_read) as read:
            carregar_aba_controle(self.path)
        self.assertTrue(any(c.kwargs.get('nrows') == 89 for c in read.call_args_list))
        self.assertEqual(self.path.read_bytes(), original)
        wb = load_workbook(self.path, data_only=False)
        self.assertEqual(wb['GERAL']['A90'].value, '=1+1')
        wb.close()

    def test_formula_com_cache_preserva_valor(self):
        def preencher(ws):
            ws['A70'], ws['B70'] = '=1+1', 'Rua A'
        self.salvar(preencher)
        with ZipFile(self.path) as arquivo:
            partes = {n: arquivo.read(n) for n in arquivo.namelist()}
        root = ET.fromstring(partes['xl/worksheets/sheet1.xml'])
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        root.find('.//m:c[@r="A70"]/m:v', ns).text = '2'
        partes['xl/worksheets/sheet1.xml'] = ET.tostring(root, encoding='utf-8')
        with ZipFile(self.path, 'w') as arquivo:
            for nome, conteudo in partes.items():
                arquivo.writestr(nome, conteudo)
        resultado, _, _ = carregar_aba_controle(self.path)
        self.assertEqual(self.limite(), 70)
        self.assertEqual(resultado.iloc[-1]['Matricula'], 2)

    def test_info_isolada_nao_infla_leitura(self):
        def preencher(ws):
            ws['A2'], ws['B2'] = 'AR0045', 'Rua A'
            ws['A12'], ws['B12'] = 'AR0046', 'Rua B'
            ws['E1048576'] = 'HM LIBERADO PARA INST'
        self.salvar(preencher)
        original = self.path.read_bytes()
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            self.assertEqual(self.limite(), 12)
            resultado, _, _ = carregar_aba_controle(self.path)
        self.assertEqual(len(resultado), 11)
        self.assertTrue(any('fora dos campos de cadastro' in str(w.message) for w in caught))
        self.assertEqual(self.path.read_bytes(), original)

    def test_somente_formatacao_e_espacos(self):
        def preencher(ws):
            ws['A2'], ws['B2'] = 'AR0045', 'Rua A'
            ws['A1048576'].number_format = '0'
            ws['B100000'] = '   '
        self.salvar(preencher)
        self.assertEqual(self.limite(), 2)
        resultado, _, _ = carregar_aba_controle(self.path)
        self.assertEqual(len(resultado), 1)

    def test_header_deslocado_e_aba_por_relacionamento(self):
        def preencher(ws):
            ws['A9'], ws['B9'] = 'AR0045', 'Rua A'
            ws.parent.create_sheet('OUTRA', 0)
        self.salvar(preencher, header=3)
        self.assertEqual(self.limite(header=2), 9)
        resultado, _, header = carregar_aba_controle(self.path)
        self.assertEqual((header, len(resultado)), (2, 6))

    def test_shared_strings_inclusive_vazia_e_zero(self):
        self.salvar(lambda ws: None)
        with ZipFile(self.path) as arquivo:
            partes = {n: arquivo.read(n) for n in arquivo.namelist()}
        ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        root = ET.fromstring(partes['xl/worksheets/sheet1.xml'])
        data = root.find('{' + ns + '}sheetData')
        for row, index in [(20, 1), (40, 0), (1000, 2)]:
            r = ET.SubElement(data, '{' + ns + '}row', r=str(row))
            c = ET.SubElement(r, '{' + ns + '}c', r='C' + str(row), t='s')
            ET.SubElement(c, '{' + ns + '}v').text = str(index)
        partes['xl/worksheets/sheet1.xml'] = ET.tostring(root, encoding='utf-8')
        partes['xl/sharedStrings.xml'] = ('<sst xmlns="' + ns + '"><si><t>0</t></si>'
                                         '<si><r><t>Parte</t></r><r><t> 2</t></r></si>'
                                         '<si><t>   </t></si></sst>').encode()
        with ZipFile(self.path, 'w') as arquivo:
            for nome, conteudo in partes.items():
                arquivo.writestr(nome, conteudo)
        self.assertEqual(self.limite(), 40)

    def test_xls_mantem_leitura_original(self):
        df = pd.DataFrame({'Matricula': ['AR0045'], 'Endereco': ['Rua A']})
        with patch('controle_utils.encontrar_linha_cabecalho', return_value=0), \
             patch('controle_utils.pd.read_excel', return_value=df) as read, \
             patch('controle_utils.determinar_ultima_linha_cadastro') as limite:
            resultado, _, _ = carregar_aba_controle('controle.xls')
        pd.testing.assert_frame_equal(resultado, df)
        self.assertNotIn('nrows', read.call_args.kwargs)
        limite.assert_not_called()

    def test_colunas_obrigatorias_ausentes(self):
        wb = Workbook()
        wb.active.title = 'GERAL'
        wb.active.append(['INFO Hidrometro'])
        wb.active['A1000'] = 'INSTALADO'
        wb.save(self.path)
        wb.close()
        with patch('controle_utils.determinar_ultima_linha_cadastro') as limite:
            self.assertIsNone(carregar_aba_controle(self.path))
        limite.assert_not_called()


if __name__ == '__main__':
    unittest.main()
